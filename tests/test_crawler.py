"""Crawl engine tests with a fake Google API (no network, no cost)."""
import openpyxl
import pytest

from backend.config import settings
from backend.crawler import engine
from backend.crawler.engine import Labeler, area_tiles, crawl_area, crawl_tile
from backend.crawler.places import PlacesClient, PlacesError
from backend.crawler.quota import QuotaManager
from backend.geo.datasets import get_country, load_areas
from backend.storage import LeadFilter
from backend.storage.sqlite_store import SQLiteLeadStore
from tests.fakes import FakeGoogle, make_place

AU = get_country("AU")
TILE = {"tile_name": "t", "low": {"latitude": -35.0, "longitude": 138.0}, "high": {"latitude": -34.0, "longitude": 139.0}}


@pytest.fixture
def quota(tmp_path):
    return QuotaManager(tmp_path / "quota.json", max_requests=1000)


def client_for(fake, quota):
    return PlacesClient(keys=["key-a", "key-b"], quota=quota, session=fake, page_delay=0)


# ---------------------------------------------------------------- places ---
def test_text_search_pages_and_geo_filter(quota):
    inside = [make_place(f"in{i}", -34.5, 138.5) for i in range(45)]
    outside = [make_place("out", -30.0, 150.0)]
    fake = FakeGoogle(places={"cafe": inside[:20] + outside + inside[20:]})
    result = client_for(fake, quota).text_search("cafe", TILE)
    assert result["requests_made"] == 3 and result["pages_fetched"] == 3
    assert len(result["results"]) == 45 and all(p["id"] != "out" for p in result["results"])
    assert quota.used == 3
    assert fake.calls[0]["json"]["locationRestriction"]["rectangle"]["low"] == TILE["low"]


def test_quota_cap_stops_before_spending(tmp_path):
    capped = QuotaManager(tmp_path / "q.json", max_requests=2)
    fake = FakeGoogle(places={"cafe": [make_place(f"p{i}", -34.5, 138.5) for i in range(60)]})
    result = client_for(fake, capped).text_search("cafe", TILE)
    assert result["quota_exceeded"] and result["requests_made"] == 2 and len(fake.calls) == 2


def test_rate_limit_rotates_key_and_retries(quota):
    fake = FakeGoogle(places={"cafe": [make_place("a", -34.5, 138.5)]}, statuses=[429])
    result = client_for(fake, quota).text_search("cafe", TILE)
    assert len(result["results"]) == 1
    assert fake.calls[0]["key"] != fake.calls[1]["key"]


def test_bad_key_raises_clear_error(quota):
    with pytest.raises(PlacesError, match="rejected the API key"):
        client_for(FakeGoogle(statuses=[403]), quota).text_search("cafe", TILE)


def test_missing_key_gives_setup_hint(quota):
    with pytest.raises(PlacesError, match="GOOGLE_API_KEYS"):
        PlacesClient(keys=[], quota=quota)


# ------------------------------------------------------------ saturation ---
def test_saturated_tile_is_split_two_levels_deep(quota):
    """Same request pattern as the original crawler: 3 + 2x3 + 4x3 = 21 requests."""
    fake = FakeGoogle(saturate=True)
    results, requests_made, quota_hit, pages = crawl_tile(client_for(fake, quota), "cafe", TILE, engine.Reporter())
    assert requests_made == 21 and pages == 3 and not quota_hit
    assert len(results) == 7 * 60  # root + 2 halves + 4 quarters, all unique


# ------------------------------------------------------------ area crawl ---
def _neighbour_point(area):
    """A point inside the area's tile but in a different council."""
    tile = area_tiles(AU, area)[0]
    councils = load_areas(AU, "lga")
    lo, hi = tile["low"], tile["high"]
    for i in range(1, 20):
        for j in range(1, 20):
            lat = lo["latitude"] + (hi["latitude"] - lo["latitude"]) * i / 20
            lon = lo["longitude"] + (hi["longitude"] - lo["longitude"]) * j / 20
            hit = councils.locate(lon, lat)
            if hit and hit.name != area.name:
                return lat, lon, hit.name
    raise AssertionError("no neighbouring council inside the tile")


def test_crawl_area_labels_leads_and_exports_only_inside(tmp_path, quota, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path / "output")
    store = SQLiteLeadStore(tmp_path / "leads.db")
    area = load_areas(AU, "lga").find("Walkerville")
    n_lat, n_lon, neighbour = _neighbour_point(area)
    inside = make_place("inside1", -34.8930, 138.6186, "Walkerville Cafe")
    both = make_place("inside2", -34.8935, 138.6190, "Walkerville Bakery Cafe")
    fake = FakeGoogle(places={
        "cafe": [inside, both, make_place("next_door", n_lat, n_lon, "Next Door Cafe")],
        "bakery": [both],
    })
    result = crawl_area(client=client_for(fake, quota), store=store, country=AU, area=area,
                        business_types=["bakery", "cafe"], types_label="2-types")

    assert result.status == "completed"
    assert (result.leads_found, result.leads_new, result.leads_outside) == (2, 2, 1)
    assert result.by_type == {"bakery": 1, "cafe": 3}

    # Every lead is stored with its real council; the neighbour is kept, not lost.
    _, leads = store.query_leads(LeadFilter())
    by_id = {l["place_id"]: l for l in leads}
    assert by_id["inside1"]["council"] == "Walkerville" and by_id["inside1"]["state"] == "South Australia"
    assert by_id["next_door"]["council"] == neighbour
    assert sorted(by_id["inside2"]["business_types"]) == ["bakery", "cafe"]

    # Excel: only businesses inside Walkerville, one row each, coordinates filled.
    wb = openpyxl.load_workbook(result.excel_file)
    rows = list(wb["Leads"].iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    assert len(data) == 2
    col = {name: i for i, name in enumerate(header)}
    assert {r[col["Business Name"]] for r in data} == {"Walkerville Cafe", "Walkerville Bakery Cafe"}
    assert all(r[col["Latitude"]] is not None and r[col["Longitude"]] is not None for r in data)
    assert all(r[col["Council"]] == "Walkerville" and r[col["State"]] == "South Australia" for r in data)
    bakery_cafe = next(r for r in data if r[col["Business Name"]] == "Walkerville Bakery Cafe")
    assert bakery_cafe[col["Business Types"]] == "bakery, cafe"
    assert "Summary" in wb.sheetnames
    assert result.excel_file.endswith(".xlsx") and "South Australia" in result.excel_file
    assert (tmp_path / "output").exists() and result.map_file.endswith("_map.html")

    runs = store.list_runs()
    assert runs[0]["status"] == "completed" and runs[0]["area"] == "Walkerville"


def test_crawl_area_stops_cleanly_at_quota(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path / "output")
    store = SQLiteLeadStore(tmp_path / "leads.db")
    area = load_areas(AU, "lga").find("Walkerville")
    capped = QuotaManager(tmp_path / "q.json", max_requests=1)
    fake = FakeGoogle(places={"bakery": [make_place("b1", -34.8930, 138.6186)]})
    result = crawl_area(client=client_for(fake, capped), store=store, country=AU, area=area,
                        business_types=["bakery", "cafe", "gym"], types_label="3-types")
    assert result.status == "quota_reached"
    assert result.leads_found == 1 and result.excel_file  # partial results still exported
    assert store.list_runs()[0]["status"] == "quota_reached"


def test_labeler_uses_coordinates_not_the_crawl_target():
    labeler = Labeler(AU)
    lead = {"latitude": -33.8568, "longitude": 151.2153}   # Sydney Opera House
    labeler.label(lead, fallback_state="Victoria")
    assert (lead["council"], lead["state"], lead["country"]) == ("Sydney", "New South Wales", "Australia")
