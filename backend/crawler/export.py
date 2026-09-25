"""
Excel and map output.

Files are organised as:
    output/<Country>/<State>/<Area>/<Area>_<types>_<YYYY-MM-DD_HHMM>.xlsx   (+ _map.html)
    output/<Country>/<State>/<City> (Nearby)/...
    output/exports/leads_<filters>_<timestamp>.xlsx                          (database exports)
"""
from __future__ import annotations

import html
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.config import settings
from backend.config.business_types import get_business_types

_BAD_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def safe_name(value: str) -> str:
    """Make a string safe to use as a Windows/Mac/Linux file or folder name."""
    cleaned = _BAD_CHARS.sub("-", str(value or "")).strip().rstrip(".")
    return cleaned or "unnamed"


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H%M")


def area_output_paths(country, area, types_label: str, dense: bool) -> dict[str, Path]:
    folder = settings.OUTPUT_DIR / safe_name(country.name) / safe_name(area.state) / safe_name(area.name)
    folder.mkdir(parents=True, exist_ok=True)
    base = f"{safe_name(area.name)}_{safe_name(types_label)}{'_dense' if dense else ''}_{_stamp()}"
    return {"excel": folder / f"{base}.xlsx", "map": folder / f"{base}_map.html"}


def city_output_paths(country, city: dict, types_label: str) -> dict[str, Path]:
    folder = settings.OUTPUT_DIR / safe_name(country.name) / safe_name(city["state"]) / f"{safe_name(city['name'])} (Nearby)"
    folder.mkdir(parents=True, exist_ok=True)
    base = f"{safe_name(city['name'])}_nearby_{safe_name(types_label)}_{_stamp()}"
    return {"excel": folder / f"{base}.xlsx", "map": folder / f"{base}_map.html"}


def export_path(label: str, create: bool = True) -> Path:
    folder = settings.OUTPUT_DIR / "exports"
    if create:
        folder.mkdir(parents=True, exist_ok=True)
    return folder / f"leads_{safe_name(label)}_{_stamp()}.xlsx"


# ------------------------------------------------------------------ Excel ---
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single")


def _categories(business_types: list[str]) -> str:
    mapping = get_business_types().type_to_category
    cats = []
    for t in business_types:
        c = mapping.get(t)
        if c and c not in cats:
            cats.append(c)
    return ", ".join(cats)


def _columns(state_label: str) -> list[tuple[str, int, object]]:
    return [
        ("Business Name", 34, lambda l: l.get("name")),
        ("Business Types", 26, lambda l: ", ".join(l.get("business_types") or [])),
        ("Category", 22, lambda l: _categories(l.get("business_types") or [])),
        ("Address", 46, lambda l: l.get("address")),
        ("Phone", 18, lambda l: l.get("phone") or l.get("phone_local")),
        ("Local Phone", 16, lambda l: l.get("phone_local")),
        ("Website", 34, lambda l: l.get("website")),
        ("Google Maps", 14, lambda l: l.get("google_maps_url")),
        ("Rating", 8, lambda l: l.get("rating")),
        ("Reviews", 9, lambda l: l.get("review_count")),
        ("Status", 18, lambda l: l.get("business_status")),
        ("Opening Hours", 48, lambda l: "; ".join(l.get("opening_hours") or [])),
        ("Council", 24, lambda l: l.get("council")),
        ("Suburb Area (SA2)", 26, lambda l: l.get("sa2")),
        (state_label, 20, lambda l: l.get("state")),
        ("Country", 12, lambda l: l.get("country")),
        ("Latitude", 11, lambda l: l.get("latitude")),
        ("Longitude", 11, lambda l: l.get("longitude")),
        ("Google Types", 30, lambda l: ", ".join(l.get("types") or [])),
        ("Place ID", 30, lambda l: l.get("place_id")),
        ("First Found", 20, lambda l: (l.get("first_seen") or "")[:16].replace("T", " ")),
        ("Last Updated", 20, lambda l: (l.get("last_seen") or "")[:16].replace("T", " ")),
    ]


def _write_leads_sheet(ws, leads: list[dict], state_label: str) -> None:
    cols = _columns(state_label)
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    link_cols = {"Website", "Google Maps"}
    for r, lead in enumerate(leads, start=2):
        for c, (title, _, getter) in enumerate(cols, start=1):
            value = getter(lead)
            if value in ("", None):
                continue
            cell = ws.cell(row=r, column=c)
            if title in link_cols and isinstance(value, str) and value.startswith("http"):
                cell.value = "Open in Maps" if title == "Google Maps" else value
                cell.hyperlink = value
                cell.font = _LINK_FONT
            else:
                cell.value = value
    for c, (_, width, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    if leads:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(leads) + 1}"


def _write_summary_sheet(ws, rows: list[tuple[str, object]], by_type: dict | None) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    bold = Font(bold=True)
    for label, value in rows:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    if by_type:
        ws.append([])
        ws.append(["Business type", "Businesses found"])
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        for btype, n in sorted(by_type.items(), key=lambda kv: (-kv[1], kv[0])):
            ws.append([btype, n])


def write_area_excel(path: Path, leads: list[dict], area, country, result, types_label: str, dense: bool,
                     city: dict | None = None, radius_m: int | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    _write_leads_sheet(ws, leads, country.state_label)
    if city:
        where = [("City (Nearby Search)", city["name"]), ("Radius", f"{radius_m / 1000:g} km"),
                 (country.state_label, city["state"])]
    else:
        level = country.levels[area.level]
        where = [("Area", area.name), ("Area type", level.label), (country.state_label, area.state),
                 ("Crawl mode", "Dense-city (SA2 neighbourhoods)" if dense else "Standard")]
    rows = where + [
        ("Country", country.name),
        ("Business types", types_label),
        ("Crawled at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Status", result.status.replace("_", " ").title()),
        ("Businesses in this file", len(leads)),
        ("New (first time found)", result.leads_new),
        ("Found in neighbouring areas", result.leads_outside if not city else "-"),
        ("Search tiles", result.tiles),
        ("Google requests used", result.requests),
        ("Estimated cost (USD)", f"${result.cost_usd:,.2f}"),
    ]
    _write_summary_sheet(wb.create_sheet("Summary"), rows, result.by_type)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_export_excel(path: Path, leads: list[dict], filters: dict, state_label: str = "State / Region") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    _write_leads_sheet(ws, leads, state_label)
    rows = [("Exported at", datetime.now().strftime("%Y-%m-%d %H:%M")), ("Businesses", len(leads))]
    rows += [(k.replace("_", " ").title(), v) for k, v in filters.items() if v not in (None, "")]
    _write_summary_sheet(wb.create_sheet("Summary"), rows, None)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ------------------------------------------------------------------- maps ---
def _popup(lead: dict) -> str:
    parts = [f"<b>{html.escape(lead.get('name') or '')}</b>"]
    for key in ("address", "phone"):
        if lead.get(key):
            parts.append(html.escape(str(lead[key])))
    if lead.get("website"):
        url = html.escape(lead["website"], quote=True)
        parts.append(f'<a href="{url}" target="_blank">{html.escape(lead["website"][:60])}</a>')
    if lead.get("business_types"):
        parts.append("<i>" + html.escape(", ".join(lead["business_types"])) + "</i>")
    return "<br>".join(parts)


def _add_leads(fmap, leads: list[dict]) -> None:
    import folium
    from folium.plugins import MarkerCluster
    cluster = MarkerCluster(name=f"Businesses ({len(leads)})").add_to(fmap)
    for lead in leads:
        if lead.get("latitude") is None:
            continue
        folium.Marker([lead["latitude"], lead["longitude"]], popup=folium.Popup(_popup(lead), max_width=320),
                      tooltip=lead.get("name")).add_to(cluster)


def write_area_map(path: Path, area, tiles: list[dict], leads: list[dict]) -> Path:
    import folium
    from shapely.geometry import mapping

    c = area.geometry.representative_point()
    fmap = folium.Map(location=[c.y, c.x], zoom_start=11, control_scale=True)
    outline = area.geometry.simplify(0.0002, preserve_topology=True)
    folium.GeoJson(mapping(outline), name=f"{area.name} boundary", tooltip=area.name,
                   style_function=lambda _: {"color": "#2E7D32", "weight": 2, "fillOpacity": 0.08}).add_to(fmap)
    tile_layer = folium.FeatureGroup(name=f"Search tiles ({len(tiles)})")
    for t in tiles:
        folium.Rectangle(bounds=[[t["low"]["latitude"], t["low"]["longitude"]], [t["high"]["latitude"], t["high"]["longitude"]]],
                         color="#EF6C00", weight=1, fill=True, fill_opacity=0.05,
                         tooltip=t.get("sub_area") or t.get("tile_name")).add_to(tile_layer)
    tile_layer.add_to(fmap)
    _add_leads(fmap, leads)
    folium.LayerControl(collapsed=False).add_to(fmap)
    minx, miny, maxx, maxy = area.geometry.bounds
    fmap.fit_bounds([[miny, minx], [maxy, maxx]])
    path.parent.mkdir(parents=True, exist_ok=True)
    fmap.save(str(path))
    return path


def write_city_map(path: Path, city: dict, radius_m: int, leads: list[dict]) -> Path:
    import folium
    fmap = folium.Map(location=[city["lat"], city["lng"]], zoom_start=12, control_scale=True)
    folium.Circle([city["lat"], city["lng"]], radius=radius_m, color="#EF6C00", fill=True, fill_opacity=0.05,
                  tooltip=f"{city['name']} search radius").add_to(fmap)
    _add_leads(fmap, leads)
    folium.LayerControl(collapsed=False).add_to(fmap)
    path.parent.mkdir(parents=True, exist_ok=True)
    fmap.save(str(path))
    return path
